"""Reviewed, atomic CSV/XLSX student onboarding. No password is accepted in files."""
import csv
import hashlib
import io
import json
import uuid
import zipfile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException
from datetime import datetime, timedelta, timezone
import jwt
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from pydantic import Field, ValidationError
from sqlalchemy import func
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from app.api.routes.teaching import admin, get_class, commit
from app.core.config import settings
from app.core.deps import AuthContext
from app.core.rate_limit import rate_limit
from app.core.security import decode_token, hash_password
from app.db.session import get_db
from app.models.academic import Enrollment, Student
from app.models.user import User, UserRole
from app.schemas.academic import Input, StudentInput
from app.services.audit_service import record_audit
from app.services.password_service import generate_temporary_password

router = APIRouter(prefix='/api/institution/student-imports', tags=['Student imports'])
COLUMNS = ['first_name', 'last_name', 'admission_number', 'login_method', 'login_id', 'email', 'phone']
MAX_BYTES = 2 * 1024 * 1024


class ImportInput(Input):
    class_id: uuid.UUID
    rows: list[dict] = Field(min_length=1, max_length=500)
    preview_token: str | None = Field(default=None, max_length=3000)


def digest(body):
    return hashlib.sha256(json.dumps(dict(class_id=str(body.class_id), rows=body.rows), sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def validate_rows(db, ctx, rows):
    errors, clean = [], []
    seen = {key: set() for key in ['admission_number', 'email', 'login_id']}
    for number, raw in enumerate(rows, 2):
        try:
            values = {k: (v.strip() if isinstance(v, str) else v) for k, v in raw.items()}
            for key in ['email', 'login_id', 'phone']:
                if values.get(key) == '': values[key] = None
            if values.get('login_method'): values['login_method'] = values['login_method'].upper()
            row = StudentInput.model_validate(values).model_dump()
            if row['email']: row['email'] = row['email'].lower()
            row_errors = []
            for key in seen:
                value = row.get(key)
                if value and value in seen[key]: row_errors.append(f'Duplicate {key} in this file')
                if value: seen[key].add(value)
            if db.query(Student.id).filter_by(institution_id=ctx.institution_id, admission_number=row['admission_number']).first():
                row_errors.append('Admission number already exists in this institution; use enrollment for an existing student')
            # Do not reveal which other institution owns a global identity.
            if row.get('email') and db.query(User.id).filter(func.lower(User.email) == row['email']).first(): row_errors.append('Email is unavailable')
            if row.get('login_id') and db.query(User.id).filter_by(student_login_id=row['login_id']).first(): row_errors.append('Student login ID is unavailable')
            errors.extend(dict(row=number, message=message) for message in row_errors)
            clean.append(row)
        except ValidationError as exc:
            errors.extend(dict(row=number, message=f"{'.'.join(map(str, e['loc']))}: {e['msg']}") for e in exc.errors())
    return clean, errors


@router.post('/parse', dependencies=[Depends(rate_limit(max_requests=10, window_seconds=60, key_prefix='import-parse'))])
async def parse_file(request: Request, filename: str = Query(max_length=200), ctx: AuthContext = Depends(admin)):
    data = bytearray()
    async for chunk in request.stream():
        data.extend(chunk)
        if len(data) > MAX_BYTES: raise HTTPException(413, 'Use a file smaller than 2 MB and at most 500 students')
    try:
        if filename.lower().endswith('.csv'):
            text = bytes(data).decode('utf-8-sig')
            reader = csv.reader(io.StringIO(text))
            records = []
            for row in reader:
                records.append(row)
                if len(records) > 501: raise ValueError('Too many rows')
        elif filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 200 or sum(f.file_size for f in archive.infolist()) > 10 * 1024 * 1024: raise ValueError('Workbook too large')
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1: raise ValueError('Use one worksheet')
                records = []
                for cells in book.active.iter_rows(max_row=502, max_col=8):
                    if any(c.data_type == 'f' for c in cells): raise ValueError('Formulas are not allowed')
                    row = [str(c.value) if c.value is not None else '' for c in cells]
                    while row and row[-1] == '': row.pop()
                    records.append(row)
                while records and not records[-1]: records.pop()
                if book.active.max_row and book.active.max_row > 501: raise ValueError('Too many rows')
                if book.active.max_column and book.active.max_column > 7: raise ValueError('Unexpected columns')
            finally: book.close()
        else: raise ValueError('Use CSV or XLSX')
        if not records or len(records) < 2 or len(records) > 501: raise ValueError('Use 1 to 500 students')
        header = [c.strip().lower() for c in records[0]]
        if len(header) != len(set(header)) or set(header) - set(COLUMNS): raise ValueError('Unexpected or duplicate column names')
        if not {'first_name', 'last_name', 'admission_number', 'login_method'}.issubset(header): raise ValueError('Missing required columns')
        rows = []
        for row in records[1:]:
            if len(row) > len(header): raise ValueError('A row has extra columns')
            rows.append(dict(zip(header, row + [''] * (len(header) - len(row)))))
        return dict(rows=rows)
    except (ValueError, UnicodeError, csv.Error, zipfile.BadZipFile, KeyError, OSError, ParseError, DefusedXmlException, InvalidFileException) as exc:
        raise HTTPException(422, f'Cannot read this student file: {exc}')


@router.post('/preview')
def preview(body: ImportInput, db: Session = Depends(get_db), ctx: AuthContext = Depends(admin)):
    cls = get_class(db, body.class_id, ctx)
    from app.services.progression import assert_group
    assert_group(db, ctx.institution_id, None, body.class_id)
    clean, errors = validate_rows(db, ctx, body.rows)
    token = None
    if not errors:
        token = jwt.encode(dict(type='student_import', sub=str(ctx.actor_user_id or ctx.user_id), user=str(ctx.user_id),
            tenant=str(ctx.institution_id), sha=digest(body), exp=datetime.now(timezone.utc) + timedelta(minutes=10)),
            settings.JWT_SECRET_KEY, algorithm=settings.JWT_ALGORITHM)
    return dict(class_name=cls.name, academic_year=cls.academic_year, students=clean, errors=errors, preview_token=token)


@router.post('/commit', dependencies=[Depends(rate_limit(max_requests=5, window_seconds=60, key_prefix='import-commit'))])
def execute(body: ImportInput, db: Session = Depends(get_db), ctx: AuthContext = Depends(admin)):
    from app.services.progression import lock_entry
    lock_entry(db, ctx.institution_id)
    try:
        token = decode_token(body.preview_token or '')
        if any(token.get(k) != v for k, v in dict(type='student_import', sub=str(ctx.actor_user_id or ctx.user_id), user=str(ctx.user_id), tenant=str(ctx.institution_id), sha=digest(body)).items()): raise ValueError()
    except (jwt.PyJWTError, ValueError): raise HTTPException(409, 'Review this file again before importing')
    get_class(db, body.class_id, ctx, lock=True)
    from app.services.progression import assert_group
    assert_group(db, ctx.institution_id, None, body.class_id)
    clean, errors = validate_rows(db, ctx, body.rows)
    if errors: raise HTTPException(409, 'Some details changed or conflict with existing accounts. Preview again; no students were imported.')
    credentials = []
    try:
        for row in clean:
            # A letter prefix keeps the generated credential literal in spreadsheets.
            password = 'S' + generate_temporary_password()
            user = User(id=uuid.uuid4(), institution_id=ctx.institution_id, username='student.' + uuid.uuid4().hex,
                first_name=row['first_name'], last_name=row['last_name'], email=row['email'], role=UserRole.STUDENT,
                student_login_method=row['login_method'], student_login_id=row['login_id'], password_hash=hash_password(password), must_change_password=True)
            db.add(user); db.flush()
            student = Student(institution_id=ctx.institution_id, user_id=user.id, admission_number=row['admission_number'], phone=row['phone'])
            db.add(student); db.flush()
            db.add(Enrollment(institution_id=ctx.institution_id, class_id=body.class_id, student_id=student.id))
            credentials.append(dict(name=f"{user.first_name} {user.last_name}", admission_number=student.admission_number,
                login_method=row['login_method'], login_identity=row['email'] if row['login_method'] == 'EMAIL' else row['login_id'], temporary_password=password))
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, 'An account detail became unavailable. No students were imported; preview the file again.')
    record_audit(db, ctx, 'students.imported', body.class_id, ctx.institution_id, f'{len(credentials)} student accounts created')
    commit(db)
    return dict(created=len(credentials), credentials=credentials)
